import re
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

# Try importing document libraries
try:
    from pypdf import PdfReader  # type: ignore
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    import docx  # type: ignore
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pptx  # type: ignore
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

class DocumentProcessor:
    @staticmethod
    def extract_text(file_path):
        """
        Extracts text from a given file path (PDF, DOCX, PPTX, TXT).
        Returns a list of dicts: [{'page_number': int, 'text': str}]
        """
        file_path = Path(file_path)
        ext = file_path.suffix.lower()
        pages = []

        if ext == ".pdf":
            if not HAS_PYPDF:
                raise ImportError("pypdf package is required for PDF parsing.")
            reader = PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append({"page_number": idx + 1, "text": text.strip()})

        elif ext in [".docx", ".doc"]:
            if not HAS_DOCX:
                raise ImportError("python-docx package is required for DOCX parsing.")
            doc = docx.Document(file_path)
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            pages.append({"page_number": 1, "text": full_text})

        elif ext in [".pptx", ".ppt"]:
            if not HAS_PPTX:
                raise ImportError("python-pptx package is required for PPTX parsing.")
            prs = pptx.Presentation(file_path)
            for idx, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    pages.append({"page_number": idx + 1, "text": "\n".join(slide_text)})

        elif ext in [".txt", ".md"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            pages.append({"page_number": 1, "text": text})

        else:
            raise ValueError(f"Unsupported file format: {ext}")

        if not pages:
            # Fallback if empty
            pages.append({"page_number": 1, "text": f"Empty document content for {file_path.name}"})

        return pages

    @staticmethod
    def clean_text(text):
        """Cleans whitespace and non-printable characters."""
        text = re.sub(r'\s+', ' ', text)
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
        return text.strip()

    @classmethod
    def process_and_chunk(cls, file_path, chunk_size=600, overlap=100):
        """
        Extracts, cleans, and chunks document content.
        Returns metadata:
        {
            "filename": str,
            "page_count": int,
            "chunks": [{'chunk_index': int, 'page_number': int, 'content': str}]
        }
        """
        file_path = Path(file_path)
        raw_pages = cls.extract_text(file_path)
        chunks = []
        chunk_idx = 0

        for page in raw_pages:
            p_num = page["page_number"]
            clean_p = cls.clean_text(page["text"])
            
            if not clean_p:
                continue

            # Word-level sliding window chunking
            words = clean_p.split(" ")
            if len(words) <= chunk_size:
                chunks.append({
                    "chunk_index": chunk_idx,
                    "page_number": p_num,
                    "content": clean_p
                })
                chunk_idx += 1
            else:
                step = chunk_size - overlap
                for i in range(0, len(words), step):
                    chunk_words = words[i : i + chunk_size]
                    if len(chunk_words) < 20 and i > 0:
                        continue # Skip trailing tiny fragments
                    chunk_str = " ".join(chunk_words)
                    chunks.append({
                        "chunk_index": chunk_idx,
                        "page_number": p_num,
                        "content": chunk_str
                    })
                    chunk_idx += 1

        return {
            "filename": file_path.name,
            "page_count": len(raw_pages),
            "chunks": chunks
        }
